import csv
import json

import openpyxl
import pytest

from table_extraction.export import TableExporter, export_to_csv, export_to_html, export_to_json
from tests.conftest import make_table
from tests.fixtures.generate import INVOICE, STAFF


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.reader(f))


def test_csv_single_table_roundtrip(tmp_path):
    out = tmp_path / "invoice.csv"
    export_to_csv([make_table(INVOICE)], out)
    assert read_csv(out) == INVOICE


def test_csv_metadata_header_for_tatr_table(tmp_path):
    out = tmp_path / "invoice.csv"
    export_to_csv([make_table(INVOICE, confidence=0.876)], out, include_metadata=True)
    rows = read_csv(out)
    assert rows[0] == ["# Table ID: 1"]
    assert rows[2] == ["# Confidence: 0.876"]
    assert rows[5:] == INVOICE


def test_json_export(tmp_path):
    out = tmp_path / "t.json"
    export_to_json([make_table(INVOICE), make_table(STAFF, table_id=2)], out, include_metadata=False)
    payload = json.loads(out.read_text())
    assert payload["extraction_info"]["total_tables"] == 2
    assert [t["data"] for t in payload["tables"]] == [INVOICE, STAFF]
    assert "metadata" not in payload["tables"][0]


def test_html_export_contains_cells(tmp_path):
    out = tmp_path / "t.html"
    export_to_html([make_table(INVOICE)], out)
    html = out.read_text()
    assert html.count("<tr>") == len(INVOICE)
    assert "<td>Widget A</td>" in html


def test_excel_export_one_sheet_per_table(tmp_path):
    out = tmp_path / "t.xlsx"
    TableExporter(tmp_path).export_tables(
        [make_table(INVOICE), make_table(STAFF, table_id=2)], "t", ["excel"]
    )
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Summary", "Table_1_Page_1", "Table_2_Page_1"]
    sheet = wb["Table_1_Page_1"]
    values = [[str(c) for c in row] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert values == INVOICE


def test_exporter_rejects_unknown_format(tmp_path):
    with pytest.raises(ValueError, match="Unsupported formats"):
        TableExporter(tmp_path).export_tables([make_table(INVOICE)], "t", ["pdf"])
